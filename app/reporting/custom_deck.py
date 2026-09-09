"""Dynamic presentation and document builder powered by Claude.

Allows Claude to design arbitrary presentations, training decks, SOPs, and reports
with AI-generated slide titles, subtitles, bullet points, and tables, compiling them
into polished PowerPoint (.pptx) or PDF (.pdf) files delivered directly to Telegram.
"""

from __future__ import annotations

import logging
import os
import tempfile
from datetime import datetime, timezone
from typing import Any, Literal, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pydantic import BaseModel, ConfigDict, Field

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.reporting.weekly_pdf import NumberedCanvas

logger = logging.getLogger("reporting.custom_deck")


class SlideItem(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(description="Slide or section title")
    subtitle: Optional[str] = Field(default=None, description="Optional subtitle or context")
    bullets: list[str] = Field(default_factory=list, description="Bullet points or key content statements")
    table_headers: Optional[list[str]] = Field(default=None, description="Optional column headers if this slide includes a table")
    table_rows: Optional[list[list[str]]] = Field(default=None, description="Optional table rows matching table_headers")


class CustomPresentationInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(description="Title of the presentation or document")
    slides: list[SlideItem] = Field(min_length=1, max_length=25, description="List of slides or sections crafted by the AI")
    format: Literal["pptx", "pdf"] = Field(
        default="pptx",
        description="Format: 'pptx' for 16:9 widescreen PowerPoint presentation, or 'pdf' for executive printable PDF document.",
    )


def build_custom_pptx(title: str, slides: list[SlideItem], output_path: str) -> str:
    """Builds a 16:9 widescreen presentation from AI-designed slides."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title Slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title
    if len(title_slide.placeholders) > 1:
        date_str = datetime.now(timezone.utc).strftime("%B %Y")
        title_slide.placeholders[1].text = f"Tara Angkor Hotel · {date_str}"

    # 2. Content Slides
    for item in slides:
        if item.table_headers and item.table_rows:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = item.title
            headers = item.table_headers
            rows = item.table_rows
            table = slide.shapes.add_table(
                len(rows) + 1, len(headers), Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.45 * (len(rows) + 1))
            ).table
            for col_idx, h in enumerate(headers):
                table.cell(0, col_idx).text = str(h)
            for r_idx, row in enumerate(rows, start=1):
                for col_idx, val in enumerate(row):
                    table.cell(r_idx, col_idx).text = str(val)
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = item.title
            body = slide.placeholders[1].text_frame
            body.clear()
            if item.subtitle:
                p = body.paragraphs[0]
                p.text = f"Context: {item.subtitle}"
                p.font.italic = True
                p.font.size = Pt(14)
                p.font.color.rgb = RGBColor(113, 128, 150)
            if item.bullets:
                start_idx = 1 if item.subtitle else 0
                for b_idx, bullet in enumerate(item.bullets):
                    p = body.add_paragraph() if (start_idx + b_idx > 0) else body.paragraphs[0]
                    p.text = bullet
                    p.font.size = Pt(16)
                    p.space_after = Pt(8)
            elif not item.subtitle:
                body.text = "Overview"

    prs.save(output_path)
    return output_path


def build_custom_pdf(title: str, slides: list[SlideItem], output_path: str) -> str:
    """Builds a styled executive PDF document from AI-designed sections."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=44,
    )

    styles = getSampleStyleSheet()

    doc_title_style = ParagraphStyle(
        "CustomDocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1A365D"),
        spaceAfter=3,
    )

    doc_subtitle_style = ParagraphStyle(
        "CustomDocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#718096"),
        spaceAfter=10,
    )

    section_heading = ParagraphStyle(
        "CustomSectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#1A365D"),
        spaceBefore=14,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "CustomDocBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#2D3748"),
    )

    bullet_style = ParagraphStyle(
        "CustomDocBullet",
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=4,
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph(title, doc_title_style))
    date_str = datetime.now(timezone.utc).strftime("%d %B %Y")
    story.append(Paragraph(f"Tara Angkor Hotel · Executive Document · {date_str}", doc_subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#D69E2E"), spaceAfter=12))

    # Sections
    for item in slides:
        section_flowables = [
            Paragraph(item.title, section_heading),
        ]
        if item.subtitle:
            section_flowables.append(Paragraph(f"<i>{item.subtitle}</i>", doc_subtitle_style))

        if item.bullets:
            for b in item.bullets:
                section_flowables.append(Paragraph(f"• {b}", bullet_style))

        if item.table_headers and item.table_rows:
            table_rows = [[Paragraph(f"<b>{h}</b>", body_style) for h in item.table_headers]]
            for row in item.table_rows:
                table_rows.append([Paragraph(str(cell), body_style) for cell in row])
            col_width = (A4[0] - 72) / max(1, len(item.table_headers))
            t = Table(table_rows, colWidths=[col_width] * len(item.table_headers))
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            section_flowables.append(Spacer(1, 4))
            section_flowables.append(t)

        section_flowables.append(Spacer(1, 8))
        story.append(KeepTogether(section_flowables))

    doc.build(story, canvasmaker=NumberedCanvas)
    return output_path


async def create_custom_presentation(args: CustomPresentationInput, context: Any) -> Any:
    """Tool handler called by Claude to build and send a custom AI-designed presentation or PDF."""
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    is_pdf = (args.format or "pptx").lower() == "pdf"
    ext = "pdf" if is_pdf else "pptx"
    path = os.path.join(tempfile.gettempdir(), f"presentation-{stamp}.{ext}")

    if is_pdf:
        build_custom_pdf(args.title, args.slides, path)
        caption = f"📄 {args.title} (PDF) — prepared for you."
        fmt_label = "PDF document"
    else:
        build_custom_pptx(args.title, args.slides, path)
        caption = f"📊 {args.title} (PowerPoint) — prepared for you."
        fmt_label = "PowerPoint presentation"

    if hasattr(context, "session") and isinstance(context.session, dict):
        context.session["pending_file"] = {"path": path, "caption": caption}

    return {
        "generated": True,
        "format": ext,
        "slides_count": len(args.slides),
        "note": f"The {fmt_label} file ({len(args.slides)} slides/sections) has been generated and is being sent to the chat now.",
    }
