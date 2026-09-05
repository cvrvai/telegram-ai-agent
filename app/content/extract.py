"""Extract bounded text for the AI provider without pretending files are trusted."""

from __future__ import annotations

import json
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

logger = logging.getLogger("content.extract")


class ContentExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedContent:
    path: str
    content_type: str
    text: str
    truncated: bool = False


class DocumentExtractor:
    """Extract bounded text from supported business documents and images."""

    SUPPORTED_TEXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}
    SUPPORTED_SPREADSHEETS = {".xlsx", ".xls"}
    SUPPORTED_IMAGES = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
    SUPPORTED_WORD = {".docx"}
    # A scanned page or photo with only a stray character or two isn't a
    # readable document -- treat it the same as "no OCR available" rather
    # than handing the model a couple of garbled letters as if it mattered.
    MIN_OCR_CHARS = 20

    def __init__(self, max_bytes: int = 10_000_000, max_chars: int = 80_000):
        self.max_bytes = max_bytes
        self.max_chars = max_chars

    def extract(self, path: str) -> ExtractedContent:
        file_path = Path(path)
        if not file_path.is_file():
            raise ContentExtractionError("The supplied file does not exist")
        if file_path.stat().st_size > self.max_bytes:
            raise ContentExtractionError("The supplied file is larger than the allowed limit")
        suffix = file_path.suffix.lower()
        if suffix in self.SUPPORTED_TEXT:
            text = file_path.read_text(encoding="utf-8", errors="replace")
            if suffix == ".json":
                try:
                    text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
                except json.JSONDecodeError:
                    pass
            return self._bounded(file_path, "text", text)
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise ContentExtractionError("PDF support requires the optional pypdf package") from exc
            try:
                text = "\n".join(page.extract_text() or "" for page in PdfReader(str(file_path)).pages)
            except Exception as exc:
                raise ContentExtractionError("The PDF could not be read") from exc
            return self._bounded(file_path, "pdf", text)
        if suffix in self.SUPPORTED_SPREADSHEETS:
            return self._extract_spreadsheet(file_path, suffix)
        if suffix in self.SUPPORTED_WORD:
            return self._extract_word(file_path)
        if suffix in self.SUPPORTED_IMAGES:
            return self._extract_image(file_path)
        raise ContentExtractionError(f"Unsupported file type: {suffix or 'unknown'}")

    def _extract_word(self, file_path: Path) -> ExtractedContent:
        try:
            import docx
        except ImportError as exc:
            raise ContentExtractionError("Word document support requires the optional python-docx package") from exc
        try:
            document = docx.Document(str(file_path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))
        except Exception as exc:
            raise ContentExtractionError("The Word document could not be read") from exc
        return self._bounded(file_path, "word", "\n".join(parts))

    def _extract_image(self, file_path: Path) -> ExtractedContent:
        """OCR the image when Tesseract is available; otherwise fall back to
        a filename placeholder the caller treats as "not readable text"."""
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return ExtractedContent(str(file_path), "image", f"Image file: {file_path.name}")
        try:
            with Image.open(file_path) as image:
                text = pytesseract.image_to_string(image)
        except Exception as exc:
            logger.warning("OCR failed for %s: %s", file_path.name, exc)
            return ExtractedContent(str(file_path), "image", f"Image file: {file_path.name}")
        text = text.strip()
        if len(text) < self.MIN_OCR_CHARS:
            return ExtractedContent(str(file_path), "image", f"Image file: {file_path.name} (no readable text found by OCR)")
        # A distinct content_type from "image" -- OCR text flows through the
        # normal text-context path instead of requiring a vision-capable
        # provider, since it's now plain extracted text like a PDF's.
        return self._bounded(file_path, "image_ocr", text)

    def _extract_spreadsheet(self, file_path: Path, suffix: str) -> ExtractedContent:
        """Convert workbook cells to bounded, sheet-labelled text for the AI."""
        try:
            if suffix == ".xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(file_path, read_only=True, data_only=True)
                chunks: list[str] = []
                for sheet in workbook.worksheets:
                    chunks.append(f"SHEET: {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = ["" if value is None else str(value) for value in row]
                        if any(values):
                            chunks.append("\t".join(values))
                workbook.close()
            else:
                import xlrd

                workbook = xlrd.open_workbook(file_path, on_demand=True)
                chunks = []
                for sheet in workbook.sheets():
                    chunks.append(f"SHEET: {sheet.name}")
                    for row in sheet.get_rows():
                        values = [cell.value if cell.value is not None else "" for cell in row]
                        rendered = [str(value) for value in values]
                        if any(rendered):
                            chunks.append("\t".join(rendered))
                workbook.release_resources()
        except ImportError as exc:
            package = "openpyxl" if suffix == ".xlsx" else "xlrd"
            raise ContentExtractionError(f"{suffix} support requires the optional {package} package") from exc
        except Exception as exc:
            raise ContentExtractionError("The spreadsheet could not be read") from exc
        return self._bounded(file_path, "spreadsheet", "\n".join(chunks))

    def _bounded(self, file_path: Path, content_type: str, text: str) -> ExtractedContent:
        truncated = len(text) > self.max_chars
        return ExtractedContent(str(file_path), content_type, text[: self.max_chars], truncated)
